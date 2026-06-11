# -*- coding: utf-8 -*-
"""Génère le rapport de fin d'études VMS (ZUM-IT) au format .docx."""

from __future__ import annotations

import datetime
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "reportss" / "Rapport_PFE_VMS_ZUM-IT.docx"
DIAGRAMS = ROOT / "diagramme_pictures"
YEAR = "2025 - 2026"
STUDENT = "Khalil ROMDHANI"
COMPANY = "ZUM-IT"
PROJECT_TITLE = (
    "Conception et réalisation d'une plateforme de gestion des visiteurs "
    "et du contrôle d'accès par QR code"
)
APP_NAME = "VMS / Zentry"


def set_doc_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    style.paragraph_format.space_after = Pt(6)
    for level, size in [(1, 16), (2, 14), (3, 13)]:
        h = doc.styles[f"Heading {level}"]
        h.font.name = "Times New Roman"
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.font.size = Pt(size)


def add_page_number_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        fld_text = OxmlElement("w:t")
        fld_text.text = "1"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_sep)
        run._r.append(fld_text)
        run._r.append(fld_end)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)


def add_header(doc: Document, text: str) -> None:
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.text = text
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.italic = True


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr)
    run._r.append(fld_char2)
    run._r.append(fld_char3)
    note = doc.add_paragraph(
        "Mettre à jour la table des matières dans Word : clic droit → « Mettre à jour les champs »."
    )
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(10)


def add_para(doc: Document, text: str, bold: bool = False, align=None) -> None:
    p = doc.add_paragraph(text)
    if bold:
        p.runs[0].bold = True
    if align is not None:
        p.alignment = align


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
    doc.add_paragraph()


def add_image_if_exists(doc: Document, path: Path, caption: str, width_cm: float = 14.0) -> None:
    if path.exists():
        doc.add_picture(str(path), width=Cm(width_cm))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(10)
    else:
        add_para(doc, f"[Figure : {caption} — insérer depuis {path.name}]")


def page_break(doc: Document) -> None:
    doc.add_page_break()


def build_report() -> Document:
    doc = Document()
    set_doc_defaults(doc)

    # ── Page de garde ─────────────────────────────────────────────
    for _ in range(6):
        doc.add_paragraph()
    t = doc.add_paragraph(YEAR)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.runs[0].bold = True
    t.runs[0].font.size = Pt(14)

    doc.add_paragraph()
    p = doc.add_paragraph("INFORMATIQUE")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(18)

    p2 = doc.add_paragraph("Développement")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.runs[0].font.size = Pt(16)

    doc.add_paragraph()
    title = doc.add_paragraph(PROJECT_TITLE)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in title.runs:
        r.bold = True
        r.font.size = Pt(14)

    doc.add_paragraph()
    sub = doc.add_paragraph(f"Application {APP_NAME} — {COMPANY}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True

    for _ in range(8):
        doc.add_paragraph()

    for label, value in [
        ("Réalisé par :", STUDENT),
        ("Encadré par (ESPRIT) :", "[Nom de l'encadrante académique]"),
        ("Encadrant entreprise :", f"[Nom — {COMPANY}]"),
    ]:
        line = doc.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = line.add_run(f"{label} ")
        r1.bold = True
        line.add_run(value)

    page_break(doc)

    # ── Dédicace ──────────────────────────────────────────────────
    doc.add_heading("Dédicace", level=1)
    add_para(
        doc,
        "Je dédie ce travail à ma famille, pour son soutien inconditionnel tout au long de mon "
        "parcours académique, ainsi qu'à tous ceux qui ont contribué, de près ou de loin, à la "
        "réussite de ce projet de fin d'études.",
    )
    page_break(doc)

    # ── Remerciements ─────────────────────────────────────────────
    doc.add_heading("Remerciements", level=1)
    add_para(
        doc,
        "Je remercie tout d'abord Dieu pour m'avoir accordé la force et la persévérance nécessaires "
        "à l'accomplissement de ce projet.",
    )
    add_para(
        doc,
        f"J'exprime ma profonde gratitude à l'entreprise {COMPANY} pour m'avoir accueilli et "
        "confié le développement du système VMS (Visitor Management System). Les échanges avec "
        "l'équipe technique ont été déterminants pour comprendre les enjeux métier et livrer une "
        "solution alignée sur les besoins réels.",
    )
    add_para(
        doc,
        "Mes remerciements s'adressent également à mon encadrante académique à l'ESPRIT, pour ses "
        "relectures, ses orientations méthodologiques et la qualité de son accompagnement.",
    )
    add_para(
        doc,
        "Enfin, je remercie les membres du jury pour l'attention qu'ils porteront à ce rapport, "
        "ainsi que mes enseignants, collègues et proches pour leurs encouragements.",
    )
    page_break(doc)

    # ── Table des matières ────────────────────────────────────────
    doc.add_heading("Table des matières", level=1)
    add_toc(doc)
    page_break(doc)

    add_header(doc, f"Rapport PFE — {APP_NAME} — {COMPANY}")
    add_page_number_footer(doc)

    # ── Introduction générale ─────────────────────────────────────
    doc.add_heading("Introduction générale", level=1)
    add_para(
        doc,
        "La gestion des visiteurs et le contrôle des accès physiques constituent un enjeu majeur "
        "pour les entreprises soucieuses de sécurité, de traçabilité et d'efficacité opérationnelle. "
        "Les procédures manuelles — registres papier, badges temporaires, validations téléphoniques — "
        "présentent des limites évidentes : lenteur, risque d'erreur, faible auditabilité et "
        "difficulté à synchroniser les informations entre accueil, sécurité et collaborateurs.",
    )
    add_para(
        doc,
        f"Dans le cadre d'un stage au sein de {COMPANY}, le projet présenté dans ce rapport vise à "
        "concevoir et réaliser une plateforme numérique complète : le Visitor Management System (VMS), "
        "commercialisé sous l'interface mobile Zentry. La solution couvre le cycle de vie d'une visite "
        "professionnelle : planification de réunion, invitation des participants, génération d'un pass "
        "numérique signé (QR code), validation à l'entrée et suivi des statuts métier.",
    )
    add_para(
        doc,
        "L'architecture retenue repose sur un backend Java Spring Boot exposant des API REST sécurisées "
        "sous le contexte /vms, une base PostgreSQL versionnée par Liquibase, et une application mobile "
        "Flutter multiplateforme consommant ces services. La sécurité s'appuie sur JWT, rôles "
        "(administrateur, employé, visiteur) et TOTP pour l'authentification visiteur.",
    )
    add_para(
        doc,
        "Ce document s'organise en sept chapitres : le contexte et la méthodologie (chapitre 1), "
        "l'analyse des besoins et la conception globale (chapitre 2), puis cinq chapitres dédiés aux "
        "sprints Agile, et une conclusion générale synthétisant les apports et perspectives.",
    )

    # ── Chapitre 1 ────────────────────────────────────────────────
    doc.add_heading("Chapitre 1 : Contexte et cadre du projet", level=1)
    doc.add_heading("Introduction", level=2)
    add_para(
        doc,
        "Ce chapitre présente l'entreprise d'accueil, le contexte métier, l'analyse de l'existant, "
        "la problématique, la solution envisagée et la méthodologie de développement adoptée.",
    )

    doc.add_heading(f"1.1 Présentation de l'entreprise {COMPANY}", level=2)
    doc.add_heading("1.1.1 Vue d'ensemble", level=3)
    add_para(
        doc,
        f"{COMPANY} est une société spécialisée dans le développement de solutions digitales et "
        "l'accompagnement technique des organisations dans la transformation de leurs processus "
        "métier. L'entreprise privilégie des architectures modernes, la qualité logicielle, la "
        "sécurité et la maintenabilité des livrables.",
    )
    doc.add_heading("1.1.2 Activités et positionnement", level=3)
    add_bullets(
        doc,
        [
            "Conception et développement d'applications métier sur mesure.",
            "Intégration de plateformes mobiles et backends REST.",
            "Accompagnement DevOps (conteneurisation, CI/CD).",
            "Sensibilisation aux bonnes pratiques de sécurité applicative.",
        ],
    )
    doc.add_heading("1.1.3 Projet confié", level=3)
    add_para(
        doc,
        "Le projet VMS a été confié dans le cadre du stage de fin d'études. Il consiste à digitaliser "
        "la gestion des visites et des réunions avec contrôle d'accès par QR code, en remplacement "
        "des flux manuels existants.",
    )

    doc.add_heading("1.2 Contexte et problématique", level=2)
    add_para(
        doc,
        "Avant la solution cible, les organisations rencontrent typiquement :",
    )
    add_bullets(
        doc,
        [
            "Un traitement manuel des visiteurs à l'accueil, source de files d'attente.",
            "Une traçabilité insuffisante des entrées et des validations.",
            "Un risque d'accès non autorisé lorsque la vérification repose sur des listes papier.",
            "Une absence de lien direct entre invitation, réunion planifiée et droit d'accès temporaire.",
        ],
    )
    add_para(
        doc,
        "La problématique centrale est donc la suivante : comment concevoir une plateforme fiable, "
        "sécurisée et utilisable sur mobile, capable d'orchestrer invitations, réunions, passes QR "
        "et validations en temps réel, tout en respectant des règles métier strictes (statuts, "
        "fenêtres horaires, rôles) ?",
    )

    doc.add_heading("1.3 Étude de l'existant", level=2)
    doc.add_heading("1.3.1 Analyse technique de la base code", level=3)
    add_para(doc, "L'audit du dépôt projet a permis d'identifier l'existant suivant :")
    add_bullets(
        doc,
        [
            "Backend Spring Boot 3.5 (Java 17) : couche web, services, repositories JPA, DTOs.",
            "Sécurité XML (vms-config.xml) : Basic Auth pour émission JWT, Bearer pour APIs protégées.",
            "PostgreSQL + migrations Liquibase (schéma profiles, meeting, office, token, confirmations).",
            "Application Flutter (zentry_staff) : écrans, controllers, services Dio, module visiteur intégré.",
            "Documentation technique : vms.md, vms-fe.md, workflow.md, setup-guide.md.",
            "Collection Postman et pipelines GitHub Actions (tests, déploiement Docker).",
        ],
    )
    doc.add_heading("1.3.2 Limites identifiées", level=3)
    add_bullets(
        doc,
        [
            "Certaines URLs encore codées en dur dans des services mobiles secondaires.",
            "Synchronisation offline avancée non implémentée (hors périmètre V1).",
            "Module entreprise/département/domaine présent côté mobile sans équivalent backend dédié.",
        ],
    )

    doc.add_heading("1.4 Objectifs et solution envisagée", level=2)
    doc.add_heading("1.4.1 Objectifs", level=3)
    add_bullets(
        doc,
        [
            "Sécuriser l'authentification employé et visiteur (JWT, TOTP).",
            "Automatiser la création et le cycle de vie des réunions.",
            "Générer et valider des QR codes signés pour l'accès physique.",
            "Offrir une expérience mobile fluide aux employés, visiteurs et agents de sécurité.",
            "Garantir la traçabilité via statuts, confirmations et journalisation applicative.",
        ],
    )
    doc.add_heading("1.4.2 Solution proposée", level=3)
    add_para(
        doc,
        "La solution VMS combine un backend REST modulaire et une application mobile Flutter. "
        "Les employés planifient des réunions et invitent visiteurs et invités internes ; les "
        "participants éligibles obtiennent un secret de réunion et signent un JWT affiché en QR ; "
        "à l'entrée, un agent scanne le code, vérifie la signature et applique les règles temporelles "
        "et de statut (CONFIRMED, IN_PROGRESS, fenêtre −15 minutes, expiration JWT).",
    )

    doc.add_heading("1.5 Méthodologie Agile / Scrum", level=2)
    doc.add_heading("1.5.1 Adoption de Scrum", level=3)
    add_para(
        doc,
        "Le projet a été mené selon une approche Agile Scrum, découpée en cinq sprints de deux à "
        "trois semaines. Chaque sprint comporte : backlog priorisé, sprint planning, développement, "
        "revue de sprint et rétrospective.",
    )
    doc.add_heading("1.5.2 Rôles", level=3)
    add_table(
        doc,
        ["Rôle", "Responsabilité"],
        [
            ["Product Owner", "Priorisation du backlog, validation métier"],
            ["Scrum Master / Encadrant", "Levée des impediments, suivi Agile"],
            ["Équipe de développement", "Implémentation backend, mobile, tests"],
        ],
    )
    doc.add_heading("1.5.3 Planification des cinq sprints", level=3)
    add_table(
        doc,
        ["Sprint", "Thème", "Livrables principaux"],
        [
            ["Sprint 1", "Authentification & setup", "JWT employé/visiteur, config projet, CI, Postman"],
            ["Sprint 2", "Fonctionnalités cœur", "Réunions CRUD, bureaux, cycle de statuts initial"],
            ["Sprint 3", "Workflows avancés", "QR (secret, scan, validation), check-in visiteur"],
            ["Sprint 4", "Notifications & intégrations", "Emails, Firebase, deep links, confirmations"],
            ["Sprint 5", "Qualité & déploiement", "Tests, Docker, documentation, optimisations"],
        ],
    )
    doc.add_heading("Conclusion du chapitre 1", level=2)
    add_para(
        doc,
        "Le cadre projet, la problématique et la méthodologie Scrum posent les fondations d'une "
        "réalisation structurée, alignée sur les besoins de ZUM-IT et sur l'architecture technique "
        "effectivement implémentée dans le dépôt VMS.",
    )

    # ── Chapitre 2 ────────────────────────────────────────────────
    doc.add_heading("Chapitre 2 : Analyse des besoins et conception", level=1)
    doc.add_heading("Introduction", level=2)

    doc.add_heading("2.1 Acteurs et besoins", level=2)
    doc.add_heading("2.1.1 Parties prenantes", level=3)
    add_table(
        doc,
        ["Acteur", "Description"],
        [
            ["Employé", "Organise des réunions, gère invitations, peut scanner des QR"],
            ["Visiteur", "Check-in, rédemption invitation, connexion TOTP, consultation réunions"],
            ["Administrateur", "Gestion employés, bureaux, supervision"],
            ["Agent sécurité", "Scan et validation des passes à l'entrée"],
        ],
    )

    doc.add_heading("2.1.2 Besoins fonctionnels", level=3)
    add_bullets(
        doc,
        [
            "Authentifier employés (email/mot de passe) et visiteurs (email/code TOTP).",
            "CRUD réunions avec organisateur, mandataire, invités, visiteurs, bureau.",
            "Transitions de statut : PENDING, CREATED, CONFIRMED, IN_PROGRESS, ENDED, CANCELED, etc.",
            "Générer un secret de réunion et un QR JWT pour les participants éligibles.",
            "Scanner et valider un QR (signature, statut, fenêtre horaire).",
            "Check-in public visiteur et rédemption de code d'invitation.",
            "Gérer les bureaux (offices) et affecter un manager.",
            "Envoyer des notifications email et push (Firebase).",
        ],
    )

    doc.add_heading("2.1.3 Besoins non fonctionnels", level=3)
    add_table(
        doc,
        ["Catégorie", "Exigence"],
        [
            ["Sécurité", "RBAC, JWT, endpoints publics limités, signature QR"],
            ["Performance", "API paginées, chargements lazy JPA"],
            ["Maintenabilité", "Couches séparées, DTOs, migrations Liquibase"],
            ["Disponibilité", "Déploiement Docker, health checks"],
            ["Portabilité", "Flutter Android/iOS"],
            ["Traçabilité", "Logs SLF4J, audit context, statuts persistés"],
        ],
    )

    doc.add_heading("2.2 Conception UML", level=2)
    doc.add_heading("2.2.1 Diagramme de cas d'utilisation global", level=3)
    add_para(
        doc,
        "Le diagramme global regroupe les packages Authentification, Administration, Réunions, "
        "Accès QR et Visiteurs. Les relations <<include>> modélisent les dépendances obligatoires "
        "(ex. connexion → émission JWT ; scan → vérification → affichage résultat). "
        "Le fichier source Mermaid est disponible sous documents/generated-diagrams/global_use_case.mmd.",
    )

    doc.add_heading("2.2.2 Diagramme de classes global", level=3)
    add_para(
        doc,
        "Le modèle domaine suit une hiérarchie User ← Employee / Visitor. Une réunion (Meeting) "
        "est liée à un bureau (Office), un organisateur, un proxy optionnel, des invités (M:N), "
        "des visiteurs (M:N) et des confirmations (composition 1-*). Les cardinalités respectent "
        "le mapping JPA du backend.",
    )
    add_image_if_exists(
        doc,
        DIAGRAMS / "class_authentication.png",
        "Figure 2.1 — Extrait du modèle d'authentification",
    )

    doc.add_heading("2.3 Architecture technique", level=2)
    doc.add_heading("2.3.1 Architecture logique", level=3)
    add_para(doc, "Trois couches backend :")
    add_bullets(
        doc,
        [
            "Présentation : MeetingWebController, EmployeeWebController, VisitorWebController, OfficeWebController, ParticipantWebController.",
            "Métier : MeetingService, EmployeeService, VisitorService, OfficeService, TokenService.",
            "Données : repositories Spring Data JPA + Querydsl, entités domain.",
        ],
    )
    add_para(doc, "Côté mobile : écrans → controllers → services API (Dio + intercepteur JWT) → modèles Dart.")

    doc.add_heading("2.3.2 Architecture physique", level=3)
    add_bullets(
        doc,
        [
            "Client mobile Flutter (smartphone/tablette).",
            "Serveur Spring Boot (context-path /vms).",
            "PostgreSQL (données métier).",
            "Firebase Cloud Messaging (notifications push).",
            "Docker Compose pour environnement local / déploiement.",
        ],
    )

    doc.add_heading("2.4 Technologies utilisées", level=2)
    add_table(
        doc,
        ["Couche", "Technologies"],
        [
            ["Backend", "Java 17, Spring Boot 3.5, Spring Security, JPA, Querydsl, Liquibase, Maven"],
            ["Base de données", "PostgreSQL (H2 en tests)"],
            ["Sécurité", "JWT, TOTP (dev.samstevens.totp), Basic Auth"],
            ["Mobile", "Flutter 3, Dart, Dio, Riverpod, qr_flutter, mobile_scanner"],
            ["Notifications", "Spring Mail, Firebase"],
            ["Qualité", "JUnit, GitHub Actions, Postman"],
        ],
    )

    doc.add_heading("2.5 Modèle de données", level=2)
    add_para(doc, "Tables principales (Liquibase) :")
    add_bullets(
        doc,
        [
            "profiles (héritage Employee/Visitor, dtype JPA)",
            "meeting, meeting_guests, meeting_visitors, meeting_confirmation",
            "office (contrainte d'unicité nom+ville+pays+société)",
            "token (invitations, récupération mot de passe, etc.)",
        ],
    )
    add_image_if_exists(
        doc,
        DIAGRAMS / "class_meetings.png",
        "Figure 2.2 — Modèle métier des réunions",
    )

    doc.add_heading("Conclusion du chapitre 2", level=2)
    add_para(
        doc,
        "L'analyse aboutit à une conception cohérente, documentée par des diagrammes UML et "
        "implémentée dans le code source analysé. Les chapitres suivants détaillent la réalisation "
        "sprint par sprint.",
    )

    # ── Helper for sprint chapters ────────────────────────────────
    def sprint_chapter(
        num: int,
        title: str,
        intro: str,
        tasks: list[str],
        use_cases: list[tuple[str, str]],
        sequences: list[str],
        realization: list[str],
        tests: list[str],
        diagram_png: str | None = None,
    ) -> None:
        doc.add_heading(f"Chapitre {num + 2} : {title}", level=1)
        doc.add_heading("Introduction", level=2)
        add_para(doc, intro)
        doc.add_heading(f"{num + 2}.1 Définition des tâches", level=2)
        add_bullets(doc, tasks)
        doc.add_heading(f"{num + 2}.2 Cas d'utilisation", level=2)
        for uc_name, uc_desc in use_cases:
            doc.add_heading(uc_name, level=3)
            add_para(doc, uc_desc)
        doc.add_heading(f"{num + 2}.3 Conception — diagrammes de séquence", level=2)
        for seq in sequences:
            add_para(doc, seq)
        if diagram_png:
            add_image_if_exists(
                doc,
                DIAGRAMS / diagram_png,
                f"Figure {num + 2}.1 — Diagramme de classes du sprint",
            )
        doc.add_heading(f"{num + 2}.4 Réalisation", level=2)
        add_bullets(doc, realization)
        doc.add_heading(f"{num + 2}.5 Tests", level=2)
        add_bullets(doc, tests)
        doc.add_heading(f"Conclusion du chapitre {num + 2}", level=2)
        add_para(
            doc,
            f"Le sprint {num} a permis de livrer les fonctionnalités prévues avec un niveau de "
            "qualité conforme aux objectifs du backlog et aux critères d'acceptation définis.",
        )

    sprint_chapter(
        1,
        "Sprint 1 — Authentification et mise en place",
        "Ce sprint établit le socle technique : configuration des projets backend et mobile, "
        "schéma de base initial, sécurité JWT et premiers écrans de connexion.",
        [
            "Initialiser le projet Spring Boot et Flutter",
            "Configurer PostgreSQL et Liquibase",
            "Implémenter POST /employees/token et POST /visitors/token",
            "Mettre en place l'intercepteur Bearer côté mobile (TokenHolder)",
            "Documenter les API (Swagger) et collection Postman",
        ],
        [
            (
                "UC01 — Se connecter (employé)",
                "L'employé saisit email et mot de passe. L'application envoie une requête "
                "POST /vms/employees/token avec en-tête Authorization: Basic. En cas de succès, "
                "le corps de réponse contient le JWT ; sinon 401.",
            ),
            (
                "UC02 — Se connecter (visiteur)",
                "Le visiteur utilise son email et un code TOTP généré à partir du secret "
                "obtenu lors du check-in. POST /vms/visitors/token retourne un JWT visiteur.",
            ),
        ],
        [
            "Séquence d'authentification employé : Mobile → API → EmployeeWebController → "
            "EmployeeService → base profiles. Branche alt : identifiants valides / invalides "
            "(voir documents/generated-diagrams/sequence_authentication.uml).",
        ],
        [
            "EmployeeWebController.login(), EmployeeService.login(), TokenService.generateJwtToken()",
            "JWTHelper et auth_interceptor.dart côté Flutter",
            "Configuration vms-config.xml (filtres Bearer, rôles admin/employee/visitor)",
        ],
        [
            "Tests Postman sur /employees/token (200 / 401)",
            "Tests unitaires services d'authentification",
            "Test manuel connexion application mobile",
        ],
        "class_authentication.png",
    )

    sprint_chapter(
        2,
        "Sprint 2 — Fonctionnalités métier cœur",
        "Ce sprint couvre la gestion des réunions et des bureaux, socle du métier VMS.",
        [
            "CRUD Meeting (POST, GET, PATCH, DELETE /meetings)",
            "Gestion Office (CRUD admin)",
            "Règles de visibilité et droits hôte/mandataire",
            "Statuts PENDING / CREATED selon créateur visiteur ou employé",
            "Écrans Flutter : liste, création, édition réunion",
        ],
        [
            (
                "UC20 — Créer une réunion",
                "Un employé (ou visiteur) crée une réunion avec sujet, créneau, bureau, "
                "organisateur, mandataire, liste d'invités et visiteurs. Le backend persiste "
                "et retourne MeetingDto.",
            ),
            (
                "UC11 — Gérer un bureau",
                "L'administrateur crée ou modifie un bureau (PATCH /offices/{id}), "
                "peut assigner un manager.",
            ),
        ],
        [
            "Séquence création réunion : POST /meetings → MeetingService.createMeeting → "
            "repository.save. Alt : créateur visiteur (PENDING) vs employé (CREATED) ; "
            "envoi emails si CREATED (sequence_meeting_creation.uml).",
            "Séquence mise à jour bureau : PATCH /offices/{id} avec alt bureau trouvé / 404.",
        ],
        [
            "MeetingService.createMeeting(), toDomain(), canAccess()",
            "MeetingWebController, OfficeWebController",
            "MeetingAddController, OfficeListController (Flutter)",
        ],
        [
            "Tests CRUD réunion via Postman",
            "Tests transitions statut initiales",
            "Tests manuels parcours création sur émulateur",
        ],
        "class_meetings.png",
    )

    sprint_chapter(
        3,
        "Sprint 3 — Workflows avancés (QR et visiteurs)",
        "Ce sprint implémente les flux à forte valeur : passes QR et parcours visiteur.",
        [
            "GET /participants/{meetingId}/secret",
            "Génération JWT QR (mid, pid, exp) côté mobile",
            "MeetingScannerController.processQRCode()",
            "POST /visitors/checkin et GET /visitors/redeem",
            "Règles MeetingScanResult (grâce 15 min, expiration)",
        ],
        [
            (
                "UC31 — Générer un QR d'accès",
                "Le participant éligible récupère le secret via l'API, signe un JWT et affiche le QR.",
            ),
            (
                "UC32 — Scanner et valider un QR",
                "L'agent scanne le code ; l'app charge la réunion, vérifie pid, récupère le secret, "
                "valide la signature et applique les règles métier.",
            ),
            (
                "UC41 — Check-in visiteur",
                "Parcours public : enregistrement ou mise à jour du profil, génération token "
                "d'invitation et envoi email avec QR TOTP.",
            ),
        ],
        [
            "sequence_qr_validation.uml : branches payload invalide, réunion absente, "
            "secret indisponible, signature OK/KO, résultat valide/trop tôt/expiré.",
            "sequence_visitor_checkin.uml : alt email employé refusé vs parcours nominal 204.",
        ],
        [
            "ParticipantWebController, MeetingService.getParticipantMeetingSecret()",
            "VisitorService.checkIn(), notifyCheckIn()",
            "qr_flutter, mobile_scanner, dart_jsonwebtoken",
        ],
        [
            "Tests scan QR valide / expiré / hors fenêtre",
            "Tests check-in et redeem sur Postman",
            "Tests manuels écran meeting_scanner.dart",
        ],
        "class_qr_access.png",
    )

    sprint_chapter(
        4,
        "Sprint 4 — Notifications, intégrations et validations",
        "Ce sprint renforce l'orchestration métier : emails, deep links, confirmations et rappels.",
        [
            "EmailTemplateService (invitations, annulation, bienvenue)",
            "DeepLinkingWebService (/links/defer, /links/redeem)",
            "GET /meetings/{id}/accept, /cancel, /pass, /start",
            "Confirmations (ACCEPTED, DECLINED, NO_RESPONSE)",
            "Firebase Cloud Messaging (fcm_token sur profiles)",
            "Tâches planifiées (rappels réunion PENDING)",
        ],
        [
            (
                "UC23 — Accepter une réunion",
                "Hôte/mandataire : PENDING→CREATED + emails visiteurs. Invité : enregistre ACCEPTED.",
            ),
            (
                "UC24 — Annuler ou refuser",
                "Hôte annule la réunion ; invité enregistre DECLINED sans annuler la réunion.",
            ),
        ],
        [
            "Notifications déclenchées après acceptation et à l'annulation (boucle visiteurs).",
            "Intégration deep link pour ouverture application depuis email.",
        ],
        [
            "MeetingService.acceptDraft(), cancelMeeting(), markConfimed(), startMeeting()",
            "MeetingRedeemerService, DeepLinkService",
            "Firebase options (firebase.json, google-services)",
        ],
        [
            "Tests emails en environnement de développement",
            "Tests endpoints accept/cancel",
            "Tests notifications push sur appareil réel",
        ],
        "class_visitors.png",
    )

    sprint_chapter(
        5,
        "Sprint 5 — Optimisation, tests et déploiement",
        "Dernier sprint : consolidation qualité, documentation et préparation déploiement.",
        [
            "Renforcement tests automatisés (GitHub Actions test.yaml)",
            "Docker Compose backend + PostgreSQL",
            "Optimisation requêtes (findByIdWithAssociations)",
            "Documentation setup-guide, workflow, diagrammes Mermaid",
            "Corrections UX et gestion erreurs API",
            "Préparation livrable PFE (rapport, démo)",
        ],
        [
            (
                "UC — Déployer l'application",
                "Packaging JAR Spring Boot, image Docker, configuration variables d'environnement.",
            ),
        ],
        [
            "Pipeline CI : build Maven, tests, rapport de couverture.",
            "docker-deploy.yaml pour déploiement conteneurisé.",
        ],
        [
            "Profils application-dev / application-test",
            "Scripts d'initialisation admin et bureaux (/opt/init)",
            "Versioning Liquibase (meeting_secret, auto_start, fcm_token)",
        ],
        [
            "Tests d'intégration backend (H2)",
            "Tests de non-régression Postman",
            "Tests manuels bout en bout : invitation → QR → scan → résultat",
            "Revue sécurité (endpoints publics, JWT TTL)",
        ],
        "class_administration.png",
    )

    # ── Conclusion générale ─────────────────────────────────────────
    doc.add_heading("Conclusion générale", level=1)
    add_para(
        doc,
        f"Le projet VMS réalisé chez {COMPANY} démontre la faisabilité d'une plateforme moderne "
        "de gestion des visiteurs et du contrôle d'accès par QR code. L'architecture en couches, "
        "la sécurité JWT/TOTP et l'application mobile Flutter offrent une réponse structurée à "
        "une problématique métier concrète.",
    )
    add_para(
        doc,
        "Les cinq sprints Agile ont permis une livraison incrémentale : authentification, réunions, "
        "QR, notifications et industrialisation. Les diagrammes UML et la documentation technique "
        "accompagnent la maintenabilité du produit.",
    )
    add_para(
        doc,
        "Les perspectives d'évolution incluent : synchronisation offline avancée, tableau de bord "
        "web administrateur, intégration contrôle d'accès physique (portiques), et analytics "
        "sur les flux de visite.",
    )

    doc.add_heading("Webographie", level=1)
    add_bullets(
        doc,
        [
            "Spring Boot Documentation — https://spring.io/projects/spring-boot",
            "Flutter Documentation — https://flutter.dev/docs",
            "PlantUML / Mermaid — modélisation UML",
            "PostgreSQL — https://www.postgresql.org/docs/",
        ],
    )

    return doc


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_report()
    doc.save(str(OUT))
    print(f"Rapport généré : {OUT}")
    print(f"Taille : {OUT.stat().st_size / 1024:.1f} Ko")


if __name__ == "__main__":
    main()
